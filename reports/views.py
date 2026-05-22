from urllib import request

from django.shortcuts import get_object_or_404

from openpyxl import Workbook
from django.contrib.auth.decorators import login_required

from django.http import HttpResponse

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)

from reportlab.lib import styles

from accounts.models import (
    TeacherProfile,
    StudentProfile
)


from forms_engine.models import (
    DynamicForm,
    FormQuestion,
    FormResponse,
    FormAnswer,
    PublicFormResponse,
    PublicFormAnswer
)
from reports.exports.identified.summary_excel import (
    export_identified_summary_excel
)

from reports.exports.identified.detailed_pdf import (
    export_identified_detailed_pdf
)



from reports.exports.anonymous.detailed_pdf import (
    export_anonymous_detailed_pdf
)


# ==========================================
# IDENTIFIED SUMMARY EXCEL
# ==========================================
import openpyxl
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.contrib.auth.decorators import login_required

# 1. IMPORT MODELS FROM YOUR FORMS ENGINE APPLICATION
# (Change 'forms_engine' to your exact forms app folder name if it's named 'forms' or 'dynamic_forms')
from forms_engine.models import DynamicForm, FormQuestion, FormResponse, FormAnswer
from accounts.models import TeacherProfile

@login_required
def identified_summary_excel(request, form_id):
    form = get_object_or_404(
        DynamicForm,
        id=form_id
    )

    # ==========================================
    # ADMIN
    # ==========================================
    if request.user.role == 'admin':
        queryset = FormResponse.objects.filter(
            form=form
        ).select_related(
            'student',
            'student__user',
            'form'
        )

    # ==========================================
    # TEACHER
    # ==========================================
    elif request.user.role == 'teacher':
        teacher_profile = get_object_or_404(
            TeacherProfile,
            user=request.user
        )

        assigned_sections = (
            teacher_profile.assigned_sections.all()
        )

        queryset = FormResponse.objects.filter(
            form=form,
            student__section__in=assigned_sections
        ).select_related(
            'student',
            'student__user',
            'form'
        )

    # ==========================================
    # OTHER USERS
    # ==========================================
    else:
        return HttpResponse('Unauthorized', status=403)


    # ==============================================================================
    # EXCEL GENERATION LOGIC
    # ==============================================================================
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Form Summary"

    # Fetch ALL dynamic questions belonging to this specific form
    questions = FormQuestion.objects.filter(form=form).order_by(
        '-is_system_field', 
        'order', 
        'id'
    )

    # Construct clean headers without time fields
    headers = ["Username", "Section", "Department"]
    for q in questions:
        headers.append(q.question)
        
    ws.append(headers)

    # Populate rows for each student submission inside the filtered queryset
    for response_obj in queryset:
        student = response_obj.student
        user = student.user if student else None
        
        # Base metadata row data (Username, Section, Department only)
        row_data = [
            user.username if user else "N/A",
            str(student.section) if student and student.section else "N/A",  # ◄── CHANGED TO THIS
            str(getattr(student, 'department', 'N/A'))
        ]

        # Sequentially map values for ALL questions without skipping anything
        for q in questions:
            # Match against your exact field definitions: 'question' and 'response'
            answer_obj = FormAnswer.objects.filter(question=q, response=response_obj).first()
            if answer_obj and answer_obj.answer:
                row_data.append(answer_obj.answer)
            else:
                row_data.append("")  # Leaves a blank space in Excel if skipped

        ws.append(row_data)

    # Package file as a seamless sheet stream back to the user's browser
    clean_title = "".join([c if c.isalnum() else "_" for c in form.title])
    
    excel_response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    excel_response["Content-Disposition"] = f'attachment; filename="{clean_title}_summary.xlsx"'
    
    wb.save(excel_response)
    return excel_response
# ==========================================
# IDENTIFIED DETAILED PDF
# ==========================================

@login_required
def identified_detailed_pdf(
    request,
    form_id
):

    form = get_object_or_404(
        DynamicForm,
        id=form_id
    )

    # ==========================================
    # ADMIN USER
    # ==========================================

    if request.user.role == 'admin':

        queryset = FormResponse.objects.filter(

            form=form

        ).select_related(

            'student',
            'student__user',
            'form'
        )

    # ==========================================
    # TEACHER USER
    # ==========================================

    elif request.user.role == 'teacher':

        teacher_profile = get_object_or_404(

            TeacherProfile,

            user=request.user
        )

        assigned_sections = (
            teacher_profile.assigned_sections.all()
        )

        queryset = FormResponse.objects.filter(

            form=form,

            student__section__in=assigned_sections

        ).select_related(

            'student',
            'student__user',
            'form'
        )

    # ==========================================
    # OTHER USERS
    # ==========================================

    else:

        return HttpResponse(
            'Unauthorized'
        )

    return export_identified_detailed_pdf(

        queryset,

        f'{form.title}_detailed'
    )


    
from django.http import HttpResponse
import csv


# ==========================================
# PUBLIC FORM EXCEL EXPORT
# ==========================================

from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from openpyxl import Workbook

# ReportLab imports needed for the PDF layout generation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# ==========================================
# 1. EXCEL EXPORT (Horizontal Rows)
# ==========================================
def public_form_excel(request, form_id):
    form = get_object_or_404(DynamicForm, id=form_id)
    
    # Optimized query prefetching answers to prevent database bottlenecks
    responses = PublicFormResponse.objects.filter(form=form).prefetch_related('publicformanswer_set')
    questions = FormQuestion.objects.filter(form=form).order_by('-is_system_field', 'order', 'id')
    
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Public Responses'

    # Append Header columns
    headers = [question.question for question in questions]
    headers.append('Submitted At')
    worksheet.append(headers)

    # Populate data side-by-side
    for response in responses:
        row = []
        answer_map = {ans.question_id: ans.answer for ans in response.publicformanswer_set.all()}

        for question in questions:
            row.append(answer_map.get(question.id, ''))

        row.append(response.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
        worksheet.append(row)

    response_file = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response_file['Content-Disposition'] = f'attachment; filename="{form.title}.xlsx"'
    workbook.save(response_file)
    return response_file


# ==========================================
# 2. PDF EXPORT (Vertical Sequential Records)
# ==========================================
def public_detailed_pdf(request, form_id):
    form = get_object_or_404(DynamicForm, id=form_id)
    responses = PublicFormResponse.objects.filter(form=form).prefetch_related('publicformanswer_set')
    questions = FormQuestion.objects.filter(form=form).order_by('-is_system_field', 'order', 'id')

    response_file = HttpResponse(content_type='application/pdf')
    response_file['Content-Disposition'] = f'attachment; filename="{form.title}.pdf"'

    doc = SimpleDocTemplate(
        response_file, 
        pagesize=letter, 
        rightMargin=40, 
        leftMargin=40, 
        topMargin=40, 
        bottomMargin=40
    )
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=22, spaceAfter=20)
    record_header_style = ParagraphStyle('RecHead', parent=styles['Heading3'], fontSize=12, spaceBefore=15, spaceAfter=10, textColor=colors.HexColor('#2563eb'))
    
    # Q_STYLE: The question label (at the top)
    q_style = ParagraphStyle('QuestText', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=14, spaceAfter=2, textColor=colors.HexColor('#0f172a'))
    
    # A_STYLE: The answer text (under the question)
    a_style = ParagraphStyle('AnsText', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14, spaceAfter=12, textColor=colors.HexColor('#334155'))

    story.append(Paragraph(f"{form.title} - Responses Report", title_style))

    for index, response in enumerate(responses, start=1):
        timestamp = response.submitted_at.strftime('%Y-%m-%d %H:%M:%S')
        story.append(Paragraph(f"Submission #{index} — Date: {timestamp}", record_header_style))
        
        answer_map = {ans.question_id: ans.answer for ans in response.publicformanswer_set.all()}

        # CORRECT ORDER: Question First, then Answer
        for question in questions:
            user_answer = answer_map.get(question.id, '—')

            # 1. ADD QUESTION FIRST
            story.append(Paragraph(f"<b>Q: {question.question}</b>", q_style))
            
            # 2. ADD ANSWER UNDER THE QUESTION
            story.append(Paragraph(f"{user_answer}", a_style))

        # Divider
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceBefore=5, spaceAfter=20))

    doc.build(story)
    return response_file
# ==========================================
# PRIVATE ANONYMOUS PDF
# ==========================================

@login_required
def private_anonymous_pdf(
    request,
    form_id
):

    form = get_object_or_404(

        DynamicForm,
        id=form_id
    )

        # ==========================================
    # ADMIN USER
    # ==========================================

    if request.user.role == 'admin':

        queryset = (
            FormResponse.objects
            .filter(
                form=form
            )
            .prefetch_related(
                'formanswer_set__question'
            )
        )

    # ==========================================
    # TEACHER USER
    # ==========================================

    elif request.user.role == 'teacher':

        teacher_profile = get_object_or_404(
            TeacherProfile,
            user=request.user
        )

        assigned_sections = (
            teacher_profile.assigned_sections.all()
        )

        queryset = (
            FormResponse.objects
            .filter(
                form=form,
                student__section__in=assigned_sections
            )
            .prefetch_related(
                'formanswer_set__question'
            )
        )

    # ==========================================
    # OTHER USERS
    # ==========================================

    else:

        return HttpResponse(
            'Unauthorized',
            status=403
        )

    return export_anonymous_detailed_pdf(

        queryset,

        f'{form.title}_private_anonymous'
    )


# ==========================================
# PRIVATE ANONYMOUS EXCEL
# =========================================
# Crucial: Import your models from your main forms engine app folder
# (Replace 'forms_engine' with 'forms' or 'dynamic_forms' if your folder name is different)
from forms_engine.models import DynamicForm, FormQuestion, FormResponse, FormAnswer
from accounts.models import TeacherProfile

# ==============================================================================
# 2. PRIVATE ANONYMOUS EXCEL DOWNLOAD (Hides ALL Identity, No Time Field)
# ==============================================================================
@login_required
def private_anonymous_excel(request, form_id):
    form = get_object_or_404(DynamicForm, id=form_id)
    responses = FormResponse.objects.filter(form=form)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Anonymous Responses'

    # Fetch questions ordered perfectly
    questions = FormQuestion.objects.filter(form=form).order_by(
        '-is_system_field', 'order', 'id'
    )

    # Headers contain ONLY questions to protect student anonymity
    headers = []
    for question in questions:
        headers.append(question.question)
    ws.append(headers)

    # Populate rows without any identity references
    for response in responses:
        row = []
        for question in questions:
            answer = FormAnswer.objects.filter(response=response, question=question).first()
            
            # Checks that object exists AND contains non-empty data safely
            if answer and answer.answer:
                row.append(answer.answer)
            else:
                row.append('')

        ws.append(row)

    clean_title = "".join([c if c.isalnum() else "_" for c in form.title])
    response_file = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response_file['Content-Disposition'] = f'attachment; filename={clean_title}_anonymous.xlsx'
    wb.save(response_file)
    return response_file